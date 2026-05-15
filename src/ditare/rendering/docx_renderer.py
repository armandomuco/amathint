"""Render canonical lesson JSON to Albanian Ditare DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ditare.generation.generate import load_lesson_json
from ditare.generation.schema import SUBJECTS, as_text, normalize_lesson, slugify
from ditare.paths import DITARE_DOCX_DIR, ensure_runtime_dirs
from ditare.validation.validate_lesson import assert_valid_lesson


def roman(num: int) -> str:
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for value, symbol in vals:
        while num >= value:
            out += symbol
            num -= value
    return out


def docx_output_path(lesson: dict, base_dir: Path = DITARE_DOCX_DIR) -> Path:
    subject_folder = SUBJECTS[lesson["subject"]]["folder"]
    filename = f"{lesson['id']}_{slugify(lesson['tema'], 50)}.docx"
    return base_dir / subject_folder / f"Klasa_{lesson['grade']}" / filename


def find_docx(lesson_id: str, base_dir: Path = DITARE_DOCX_DIR) -> Path | None:
    matches = sorted(base_dir.rglob(f"{lesson_id}_*.docx"))
    return matches[0] if matches else None


class DitareDocxRenderer:
    def render_to_path(self, lesson: dict, output_path: str | Path | None = None, force: bool = False) -> Path:
        ensure_runtime_dirs()
        lesson = normalize_lesson(lesson)
        assert_valid_lesson(lesson)
        path = Path(output_path) if output_path else docx_output_path(lesson)
        if path.exists() and not force:
            return path
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)

        self._add_title_and_date(doc)
        table = doc.add_table(rows=7, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, width in enumerate([Cm(4), Cm(4), Cm(4), Cm(4)]):
            for cell in table.columns[idx].cells:
                cell.width = width

        self._fill_header(table, lesson)
        self._fill_topic(table, lesson)
        self._fill_results(table, lesson)
        self._fill_resources(table, lesson)
        self._fill_methodology_header(table)
        self._fill_body(table, lesson)
        self._fill_homework(table, lesson)

        doc.save(path)
        return path

    def render_lesson_id(self, lesson_id: str, force: bool = False) -> Path:
        lesson = load_lesson_json(lesson_id)
        return self.render_to_path(lesson, force=force)

    def _run(self, paragraph, text: str, bold: bool = False):
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(10)
        return run

    def _shade(self, cell, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        tc_pr.append(shading)

    def _remove_borders(self, table) -> None:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "nil")
            borders.append(border)
        tbl_pr.append(borders)

    def _add_title_and_date(self, doc: Document) -> None:
        header = doc.add_table(rows=1, cols=2)
        header.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._remove_borders(header)

        title_cell = header.cell(0, 0)
        title_cell.width = Cm(11)
        title = title_cell.paragraphs[0]
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = self._run(title, "PLANIFIKIMI I ORËS MËSIMORE", bold=True)
        title_run.font.name = "Times New Roman"
        title_run.font.size = Pt(16)

        date_cell = header.cell(0, 1)
        date_cell.width = Cm(5)
        date = date_cell.paragraphs[0]
        date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        date_run = self._run(date, "Data ___ / ___ / _____", bold=True)
        date_run.font.name = "Times New Roman"
        date_run.font.size = Pt(12)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    def _merge(self, table, row: int, start: int, end: int):
        return table.cell(row, start).merge(table.cell(row, end))

    def _fill_header(self, table, lesson: dict) -> None:
        subject_info = SUBJECTS[lesson["subject"]]
        values = [
            f"Fusha {subject_info['fusha']}",
            f"Lënda {lesson['subject_full']}",
            f"Shkalla {lesson['shkalla']}",
            f"Klasa {roman(lesson['grade'])}",
        ]
        for idx, value in enumerate(values):
            cell = table.rows[0].cells[idx]
            self._shade(cell, "E6E6E6")
            self._run(cell.paragraphs[0], value, bold=True)

    def _fill_topic(self, table, lesson: dict) -> None:
        topic_cell = self._merge(table, 1, 0, 1)
        p = topic_cell.paragraphs[0]
        self._run(p, f"Tema mësimore nr. {lesson['lesson_number']}: ", bold=True)
        self._run(p, lesson["tema"])

        situation_cell = self._merge(table, 1, 2, 3)
        p = situation_cell.paragraphs[0]
        self._run(p, "Situata e të nxënit:\n", bold=True)
        self._run(p, lesson["situata"])

    def _fill_results(self, table, lesson: dict) -> None:
        results_cell = self._merge(table, 2, 0, 1)
        p = results_cell.paragraphs[0]
        self._run(p, "Rezultatet e të nxënit të kompetencave të fushës sipas temës mësimore:\n", bold=True)
        for result in lesson["rezultatet"]:
            self._run(p, f"• {result}\n")

        keywords_cell = self._merge(table, 2, 2, 3)
        p = keywords_cell.paragraphs[0]
        self._run(p, "Fjalët kyçe:\n", bold=True)
        self._run(p, "\n".join(lesson["fjalet_kyce"]))

    def _fill_resources(self, table, lesson: dict) -> None:
        p = table.rows[3].cells[0].paragraphs[0]
        self._run(p, "Burimet:\n", bold=True)
        self._run(p, "\n".join(lesson["burimet"]))

        link_cell = self._merge(table, 3, 1, 3)
        p = link_cell.paragraphs[0]
        self._run(p, "Lidhja me fushat e tjera ose me temat ndërkurrikulare:\n", bold=True)
        self._run(p, lesson["lidhja"])

    def _fill_methodology_header(self, table) -> None:
        cell = self._merge(table, 4, 0, 3)
        self._shade(cell, "D9D9D9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, "Metodologjia e realizimit të temës mësimore", bold=True)

    def _fill_body(self, table, lesson: dict) -> None:
        cell = self._merge(table, 5, 0, 3)
        p = cell.paragraphs[0]
        self._run(p, "Organizimi i orës mësimore:\n", bold=True)
        labels = [
            ("qellimi", "Qëllimi"),
            ("evokim", "Evokimi/Hyrja"),
            ("realizim_kuptimi", "Realizimi i kuptimit/Zhvillimi"),
            ("praktike", "Praktika/Ushtrime"),
            ("reflektim", "Mbyllja/Reflektimi"),
            ("vleresim", "Vlerësimi"),
            ("nxenes_ak", "Nxënës me AK dhe vështirësi në të nxënë"),
        ]
        for key, label in labels:
            value = as_text(lesson["metodologjia"].get(key))
            if value:
                self._run(p, f"\n{label}:\n", bold=True)
                self._run(p, f"{value}\n")

    def _fill_homework(self, table, lesson: dict) -> None:
        cell = self._merge(table, 6, 0, 3)
        p = cell.paragraphs[0]
        self._run(p, "Detyrat e dhëna për punë të pavarur:\n", bold=True)
        detyrat = lesson["detyrat"]
        if detyrat.get("baze"):
            self._run(p, f"Bazë: {detyrat['baze']}\n")
        if detyrat.get("krijuese"):
            self._run(p, f"Krijuese: {detyrat['krijuese']}\n")
        if detyrat.get("shtese"):
            self._run(p, f"Shtesë: {detyrat['shtese']}\n")


def render_lessons(lessons: list[dict], force: bool = False) -> list[Path]:
    renderer = DitareDocxRenderer()
    return [renderer.render_to_path(lesson, force=force) for lesson in lessons]
