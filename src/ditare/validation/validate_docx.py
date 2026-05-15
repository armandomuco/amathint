"""Validation for rendered Amathint DOCX files."""

from __future__ import annotations

from pathlib import Path

from docx import Document


def _row_text(table, row_idx: int) -> str:
    return " ".join(" ".join(cell.text.split()) for cell in table.rows[row_idx].cells)


def validate_docx(path: str | Path, min_body_chars: int = 500) -> list[str]:
    errors: list[str] = []
    path = Path(path)
    if not path.exists():
        return [f"file does not exist: {path}"]
    try:
        doc = Document(path)
    except Exception as exc:
        return [f"cannot open docx: {exc}"]
    if not doc.tables:
        return ["document has no Amathint/Ditare table"]
    table = doc.tables[0]
    if len(table.rows) < 7 or len(table.columns) < 4:
        errors.append("Amathint/Ditare table must have at least 7 rows and 4 columns")
        return errors

    row0 = _row_text(table, 0)
    row2 = _row_text(table, 2)
    row3 = _row_text(table, 3)
    row5 = _row_text(table, 5)
    row6 = _row_text(table, 6)

    if "Fusha" not in row0 or "Lënda" not in row0 or "Klasa" not in row0:
        errors.append("row 0 must contain subject and class headers")
    if row0.strip().endswith("Klasa"):
        errors.append("class value is blank")
    if len(row2.replace("Fjalët kyçe:", "").strip()) < 20:
        errors.append("keywords/results row is incomplete")
    if len(row3.replace("Lidhja me fushat e tjera ose me temat ndërkurrikulare:", "").strip()) < 20:
        errors.append("resources/cross-curricular row is incomplete")
    if len(row5) < min_body_chars:
        errors.append(f"lesson body must be at least {min_body_chars} characters")
    if len(row6.replace("Detyrat e dhëna për punë të pavarur:", "").strip()) < 20:
        errors.append("homework row is incomplete")
    return errors
