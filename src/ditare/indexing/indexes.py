"""Build searchable JSON and Excel indexes."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ditare.curriculum.catalog import load_catalog
from ditare.generation.generate import find_lesson_json
from ditare.paths import DEFAULT_INDEX_JSON, DEFAULT_INDEX_XLSX, INDEXES_DIR, ensure_runtime_dirs
from ditare.rendering.docx_renderer import find_docx
from ditare.validation.validate_docx import validate_docx
from ditare.validation.validate_lesson import validate_lesson


def _load_json(path: Path) -> dict | None:
    if not path:
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def build_index(
    catalog_path: str | Path | None = None,
    output_json: str | Path = DEFAULT_INDEX_JSON,
    output_xlsx: str | Path = DEFAULT_INDEX_XLSX,
    subject: str | None = None,
) -> dict:
    ensure_runtime_dirs()
    catalog = load_catalog(catalog_path) if catalog_path else load_catalog()
    rows = []
    for topic in catalog.get("topics", []):
        if subject and topic.get("subject") != subject:
            continue
        json_path = find_lesson_json(topic["id"])
        docx_path = find_docx(topic["id"])
        lesson = _load_json(json_path) if json_path else None
        json_errors = validate_lesson(lesson) if lesson else ["canonical JSON missing"]
        docx_errors = validate_docx(docx_path) if docx_path else ["DOCX missing"]
        rows.append(
            {
                "id": topic["id"],
                "subject": topic["subject"],
                "subject_full": topic["subject_full"],
                "grade": topic["grade"],
                "lesson_number": topic["lesson_number"],
                "tema": topic["tema"],
                "tematika": topic["tematika"],
                "month": topic["month"],
                "json_path": str(json_path) if json_path else "",
                "docx_path": str(docx_path) if docx_path else "",
                "json_valid": not json_errors,
                "docx_valid": not docx_errors,
                "json_errors": json_errors,
                "docx_errors": docx_errors,
            }
        )

    index = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "total_topics": len(rows),
        "generated_json": sum(1 for row in rows if row["json_path"]),
        "rendered_docx": sum(1 for row in rows if row["docx_path"]),
        "valid_json": sum(1 for row in rows if row["json_valid"]),
        "valid_docx": sum(1 for row in rows if row["docx_valid"]),
        "lessons": rows,
    }

    output_json = Path(output_json)
    output_json.parent.mkdir(parents=True, exist_ok=True)
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)

    write_xlsx(index, output_xlsx)
    return index


def write_xlsx(index: dict, output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Amathint Index"
    headers = [
        "ID",
        "Subject",
        "Grade",
        "Lesson",
        "Tema",
        "Tematika",
        "Month",
        "JSON Valid",
        "DOCX Valid",
        "JSON Path",
        "DOCX Path",
        "Errors",
    ]
    ws.append(headers)
    for lesson in index["lessons"]:
        ws.append(
            [
                lesson["id"],
                lesson["subject_full"],
                lesson["grade"],
                lesson["lesson_number"],
                lesson["tema"],
                lesson["tematika"],
                lesson["month"],
                lesson["json_valid"],
                lesson["docx_valid"],
                lesson["json_path"],
                lesson["docx_path"],
                "; ".join(lesson["json_errors"] + lesson["docx_errors"]),
            ]
        )

    fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        ws.column_dimensions[column[0].column_letter].width = width

    wb.save(output_path)
    return output_path


def index_status() -> dict:
    return {
        "index_dir": str(INDEXES_DIR),
        "json_exists": DEFAULT_INDEX_JSON.exists(),
        "xlsx_exists": DEFAULT_INDEX_XLSX.exists(),
        "json_path": str(DEFAULT_INDEX_JSON),
        "xlsx_path": str(DEFAULT_INDEX_XLSX),
    }
